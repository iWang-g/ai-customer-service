from __future__ import annotations

import io
import json
import re
from dataclasses import dataclass
from typing import Any


SUPPORTED_EXTENSIONS = {".txt", ".md", ".csv", ".json", ".html", ".docx", ".pdf", ".xlsx", ".xlsm"}
CHUNK_STRATEGY_VERSION = "structured-v2"


@dataclass(frozen=True)
class ParsedDocument:
    text: str
    file_type: str


@dataclass(frozen=True)
class TextChunk:
    index: int
    title_path: str
    content: str
    chunk_type: str = "prose"
    metadata: dict[str, Any] | None = None


def _decode_text(data: bytes) -> str:
    for encoding in ("utf-8-sig", "gb18030", "utf-16"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    return data.decode("utf-8", errors="replace")


def parse_document(filename: str, data: bytes) -> ParsedDocument:
    suffix = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if suffix not in SUPPORTED_EXTENSIONS:
        raise ValueError(f"unsupported file type: {suffix or 'unknown'}")
    if suffix in {".txt", ".md", ".csv", ".json", ".html"}:
        text = _decode_text(data)
        if suffix == ".json":
            try:
                text = json.dumps(json.loads(text), ensure_ascii=False, indent=2)
            except ValueError:
                pass
        elif suffix == ".html":
            text = re.sub(r"(?is)<(?:script|style)[^>]*>.*?</(?:script|style)>", "", text)
            text = re.sub(r"(?i)</?(?:h([1-6])|p|div|li|tr|br)[^>]*>", "\n", text)
            text = re.sub(r"<[^>]+>", "", text)
    elif suffix == ".docx":
        try:
            from docx import Document
        except ImportError as exc:
            raise ValueError("DOCX 解析依赖未安装，请安装 python-docx") from exc
        document = Document(io.BytesIO(data))
        paragraphs: list[str] = []
        for paragraph in document.paragraphs:
            value = paragraph.text.strip()
            if not value:
                continue
            style_name = str(paragraph.style.name or "") if paragraph.style else ""
            heading_match = re.search(r"(?:heading|标题)\s*([1-6])", style_name, re.IGNORECASE)
            paragraphs.append(f"{'#' * int(heading_match.group(1))} {value}" if heading_match else value)
        text = "\n\n".join(paragraphs)
        for table in document.tables:
            rows = [" | ".join(cell.text.strip() for cell in row.cells) for row in table.rows]
            text += "\n\n" + "\n".join(row for row in rows if row.strip())
    elif suffix == ".pdf":
        try:
            from pypdf import PdfReader
        except ImportError as exc:
            raise ValueError("PDF 解析依赖未安装，请安装 pypdf") from exc
        reader = PdfReader(io.BytesIO(data))
        pages = []
        for page_number, page in enumerate(reader.pages, start=1):
            page_text = (page.extract_text() or "").strip()
            if page_text:
                pages.append(f"# 第 {page_number} 页\n\n{page_text}")
        text = "\n\n".join(pages).strip()
    else:
        try:
            from openpyxl import load_workbook
        except ImportError as exc:
            raise ValueError("XLSX 解析依赖未安装，请安装 openpyxl") from exc
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
        sections: list[str] = []
        for sheet in workbook.worksheets:
            sections.append(f"# {sheet.title}")
            for row in sheet.iter_rows(values_only=True):
                values = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if values:
                    sections.append(" | ".join(values))
        text = "\n".join(sections)
    text = re.sub(r"\r\n?", "\n", text).strip()
    if not text:
        raise ValueError("文档没有可解析的文本内容")
    return ParsedDocument(text=text, file_type=suffix.lstrip("."))


def _heading(line: str) -> tuple[int, str] | None:
    match = re.match(r"^(#{1,6})\s+(.+?)\s*$", line.strip())
    if not match:
        return None
    return len(match.group(1)), match.group(2).strip()


def _block_type(block: str) -> str:
    lines = [line.strip() for line in block.splitlines() if line.strip()]
    if lines and all("|" in line or "\t" in line for line in lines):
        return "table"
    if lines and all(re.match(r"^(?:[-*+]\s+|\d+[.)、]\s*)", line) for line in lines):
        return "list"
    if lines and all(re.match(r"^[^：:\n]{1,40}[：:]\s*.+", line) for line in lines):
        return "parameters"
    return "prose"


def _semantic_units(block: str, chunk_type: str) -> list[str]:
    if chunk_type in {"table", "list", "parameters"}:
        return [line.strip() for line in block.splitlines() if line.strip()]
    normalized = re.sub(r"[ \t]+", " ", block.strip())
    units = [item.strip() for item in re.split(r"(?<=[。！？!?；;])\s*|\n+", normalized) if item.strip()]
    return units or [normalized]


def _split_oversized_unit(unit: str, max_chars: int) -> list[str]:
    pieces: list[str] = []
    remaining = unit.strip()
    while len(remaining) > max_chars:
        boundary = max(
            remaining.rfind(marker, 0, max_chars + 1)
            for marker in ("。", "！", "？", "；", ";", "，", ",", " ")
        )
        end = boundary + 1 if boundary >= max_chars // 2 else max_chars
        pieces.append(remaining[:end].strip())
        remaining = remaining[end:].strip()
    if remaining:
        pieces.append(remaining)
    return pieces


def chunk_text(
    text: str,
    target_chars: int = 520,
    max_chars: int = 800,
    overlap_chars: int = 80,
) -> list[TextChunk]:
    cleaned = re.sub(r"\n{3,}", "\n\n", text.replace("\u00a0", " ")).strip()
    if not cleaned:
        return []
    chunks: list[TextChunk] = []
    heading_stack: list[str] = []
    current_units: list[str] = []
    current_type = "prose"
    current_title_path = ""
    overlap_only = False

    def flush(*, keep_overlap: bool = False) -> None:
        nonlocal current_units, overlap_only
        content = "\n".join(current_units).strip()
        if not content:
            current_units = []
            overlap_only = False
            return
        if overlap_only:
            current_units = []
            overlap_only = False
            return
        chunks.append(TextChunk(
            index=len(chunks),
            title_path=current_title_path,
            content=content,
            chunk_type=current_type,
            metadata={"strategy_version": CHUNK_STRATEGY_VERSION},
        ))
        previous_units = current_units
        current_units = []
        overlap_only = False
        if keep_overlap and current_type == "prose" and overlap_chars > 0:
            overlap: list[str] = []
            length = 0
            for unit in reversed(previous_units):
                if overlap and length + len(unit) > overlap_chars:
                    break
                overlap.insert(0, unit)
                length += len(unit)
            if overlap and length < len(content):
                current_units = overlap
                overlap_only = True

    for raw_block in re.split(r"\n\s*\n", cleaned):
        block = raw_block.strip()
        if not block:
            continue
        heading = _heading(block) if "\n" not in block else None
        if heading:
            flush()
            level, title = heading
            heading_stack[level - 1:] = [title]
            current_title_path = " > ".join(heading_stack)
            current_units = []
            overlap_only = False
            continue
        block_type = _block_type(block)
        if current_units and block_type != current_type:
            flush()
            current_units = []
        current_type = block_type
        current_title_path = " > ".join(heading_stack)
        for unit in _semantic_units(block, block_type):
            for piece in _split_oversized_unit(unit, max_chars):
                current_length = sum(len(item) for item in current_units) + len(current_units)
                if current_units and current_length >= target_chars:
                    flush(keep_overlap=True)
                    current_length = sum(len(item) for item in current_units) + len(current_units)
                projected = current_length + len(piece)
                if current_units and projected > max_chars:
                    flush(keep_overlap=True)
                    if sum(len(item) for item in current_units) + len(current_units) + len(piece) > max_chars:
                        current_units = []
                        overlap_only = False
                current_units.append(piece)
                overlap_only = False
    flush()
    return chunks
