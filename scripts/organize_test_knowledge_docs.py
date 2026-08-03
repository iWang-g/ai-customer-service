from __future__ import annotations

import json
import re
import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION_START
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


ROOT = Path(__file__).resolve().parents[1]
SOURCE_DIR = ROOT / "test-knowledge-docs"
OUTPUT_DIR = SOURCE_DIR / "organized"
EXTRACT_DIR = OUTPUT_DIR / "_extracted"


@dataclass(frozen=True)
class Block:
    source: str
    kind: str
    text: str


def normalize_text(value: str) -> str:
    value = value.replace("\u3000", " ")
    value = re.sub(r"[ \t]+", " ", value)
    value = re.sub(r"\n{3,}", "\n\n", value)
    return value.strip()


SPLIT_MARKERS = [
    "店铺：",
    "售前",
    "售后部分",
    "售后问题及解决方法",
    "以下为产品售后的解决方案：",
    "基础知识：",
    "销售策略：",
    "销售核心话术",
    "关联销售",
    "主要产品A：",
    "主要产品B：",
    "产品C：",
    "主要产品D：",
    "分腿抱枕",
    "等身抱枕：",
    "抱枕枕芯有",
    "枕套分为几种不同的材质：",
    "关于定制图：",
    "图库系统：",
    "一代义乳：",
    "二代义乳：",
    "三代义乳：",
    "半身起毛套：",
    "全身起毛套：",
    "分腿起毛套：",
    "未发货退款：",
    "已发货退款：",
    "退货退款：",
    "包裹破损：",
    "商品瑕疵：",
    "厂家发错：",
    "快递问题：",
    "退货处理：",
    "商品丢件：",
    "枕芯枕套瑕疵",
    "枕套枕芯发错货：",
    "枕套图案发错：",
    "客人退货退款七天内：",
    "客人退货退款七天后：",
    "义臀售后：",
    "义臀瑕疵：",
    "退货退款处理：",
]


def split_logical_blocks(text: str) -> list[str]:
    text = normalize_text(text)
    if not text:
        return []
    if len(text) < 1200:
        return [text]

    marked = text
    for marker in sorted(SPLIT_MARKERS, key=len, reverse=True):
        marked = marked.replace(marker, f"\n\n{marker}")
    marked = re.sub(r"\s+(\d+[.．、])", r"\n\1", marked)
    chunks = [normalize_text(chunk) for chunk in re.split(r"\n{2,}", marked) if normalize_text(chunk)]

    result: list[str] = []
    for chunk in chunks:
        if len(chunk) <= 1600:
            result.append(chunk)
            continue
        sentences = re.split(r"(?<=[。！？!?])", chunk)
        buffer = ""
        for sentence in sentences:
            sentence = normalize_text(sentence)
            if not sentence:
                continue
            if len(buffer) + len(sentence) > 1000 and buffer:
                result.append(buffer)
                buffer = sentence
            else:
                buffer = normalize_text(f"{buffer} {sentence}")
        if buffer:
            result.append(buffer)
    return result or [text]


def iter_doc_blocks(path: Path) -> Iterable[Block]:
    doc = Document(path)
    for paragraph in doc.paragraphs:
        text = normalize_text(paragraph.text)
        for idx, chunk in enumerate(split_logical_blocks(text), start=1):
            kind = "paragraph" if idx == 1 else f"paragraph-split-{idx}"
            yield Block(path.name, kind, chunk)
    for table_index, table in enumerate(doc.tables, start=1):
        rows: list[str] = []
        for row in table.rows:
            cells = [normalize_text(cell.text) for cell in row.cells]
            if any(cells):
                rows.append(" | ".join(cells))
        if rows:
            for idx, chunk in enumerate(split_logical_blocks("\n".join(rows)), start=1):
                kind = f"table-{table_index}" if idx == 1 else f"table-{table_index}-split-{idx}"
                yield Block(path.name, kind, chunk)


def set_font(run, name: str = "Microsoft YaHei", size: float | None = None, bold: bool | None = None):
    run.font.name = name
    run._element.rPr.rFonts.set(qn("w:ascii"), name)
    run._element.rPr.rFonts.set(qn("w:hAnsi"), name)
    run._element.rPr.rFonts.set(qn("w:eastAsia"), name)
    if size is not None:
        run.font.size = Pt(size)
    if bold is not None:
        run.bold = bold


def configure_doc(doc: Document, title: str, subtitle: str):
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(0.82)
    section.right_margin = Inches(0.82)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Microsoft YaHei"
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(5)
    normal.paragraph_format.line_spacing = 1.2

    for style_name, size, color in [
        ("Heading 1", 16, RGBColor(46, 116, 181)),
        ("Heading 2", 13, RGBColor(46, 116, 181)),
        ("Heading 3", 11.5, RGBColor(31, 77, 120)),
    ]:
        style = styles[style_name]
        style.font.name = "Microsoft YaHei"
        style._element.rPr.rFonts.set(qn("w:eastAsia"), "Microsoft YaHei")
        style.font.size = Pt(size)
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(8)
        style.paragraph_format.space_after = Pt(4)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    r = p.add_run(title)
    set_font(r, size=20, bold=True)
    r.font.color.rgb = RGBColor(11, 37, 69)

    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(12)
    r = p.add_run(subtitle)
    set_font(r, size=9.5)
    r.font.color.rgb = RGBColor(85, 85, 85)


def add_toc(doc: Document, items: list[tuple[str, str]]):
    doc.add_heading("后台录入定位", level=1)
    table = doc.add_table(rows=1, cols=2)
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    set_table_widths(table, [1800, 7560])
    hdr = table.rows[0].cells
    hdr[0].text = "后台字段"
    hdr[1].text = "整理后的承载内容"
    for left, right in items:
        cells = table.add_row().cells
        cells[0].text = left
        cells[1].text = right
    format_table(table)


def add_source_note(doc: Document, sources: Iterable[str]):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(4)
    p.paragraph_format.space_after = Pt(10)
    r = p.add_run("来源文档：")
    set_font(r, bold=True)
    p.add_run("、".join(sources))


def set_table_widths(table, widths: list[int]):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    tbl_w = tbl_pr.find(qn("w:tblW"))
    if tbl_w is None:
        tbl_w = OxmlElement("w:tblW")
        tbl_pr.insert(0, tbl_w)
    tbl_w.set(qn("w:type"), "dxa")
    tbl_w.set(qn("w:w"), str(sum(widths)))
    tbl_ind = OxmlElement("w:tblInd")
    tbl_ind.set(qn("w:w"), "120")
    tbl_ind.set(qn("w:type"), "dxa")
    tbl_pr.append(tbl_ind)

    tbl_grid = tbl.tblGrid
    for child in list(tbl_grid):
        tbl_grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol")
        col.set(qn("w:w"), str(width))
        tbl_grid.append(col)

    for row in table.rows:
        for idx, cell in enumerate(row.cells):
            tc_pr = cell._tc.get_or_add_tcPr()
            tc_w = tc_pr.tcW
            tc_w.set(qn("w:type"), "dxa")
            tc_w.set(qn("w:w"), str(widths[idx]))


def shade_cell(cell, fill: str):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def format_table(table):
    for row_index, row in enumerate(table.rows):
        for cell in row.cells:
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            for paragraph in cell.paragraphs:
                paragraph.paragraph_format.space_after = Pt(2)
                for run in paragraph.runs:
                    set_font(run, size=9.5, bold=(row_index == 0))
            if row_index == 0:
                shade_cell(cell, "E8EEF5")


def add_bullets(doc: Document, values: Iterable[str]):
    for value in values:
        text = normalize_text(value)
        if text:
            p = doc.add_paragraph(style="List Bullet")
            p.add_run(text)


def keyword_candidates(text: str) -> list[str]:
    keywords: list[str] = []
    patterns = [
        r"邮箱",
        r"邮件",
        r"链接",
        r"售前",
        r"售后",
        r"发货",
        r"物流",
        r"退货",
        r"退款",
        r"换货",
        r"优惠",
        r"价格",
        r"客服",
        r"二次元",
        r"极霸猫",
        r"宅小兔",
    ]
    for pattern in patterns:
        if re.search(pattern, text) and pattern not in keywords:
            keywords.append(pattern)
    return keywords[:8]


def split_question_answer(blocks: list[Block]) -> list[dict[str, str | list[str] | int | bool]]:
    entries: list[dict[str, str | list[str] | int | bool]] = []
    seen: set[tuple[str, str]] = set()

    def add_entry(category: str, question: str, answer: str, source: str, weight: int = 10):
        question = normalize_text(question)
        answer = normalize_text(answer)
        if not question or not answer:
            return
        key = (question, answer[:120])
        if key in seen:
            return
        seen.add(key)
        entries.append(
            {
                "category": category,
                "question": question,
                "keywords": keyword_candidates(question + answer),
                "answer": answer,
                "image_url": "",
                "weight": weight,
                "enabled": True,
                "source": source,
            }
        )

    for block in blocks:
        text = block.text
        lines = [normalize_text(line) for line in text.splitlines() if normalize_text(line)]
        first_line_is_question = bool(re.match(r"^(Q|问|问题)[:：\s]", lines[0], flags=re.I)) if lines else False
        first_line_is_question = first_line_is_question or (bool(lines) and len(lines[0]) <= 120 and lines[0].endswith(("？", "?")))
        if len(lines) >= 2 and first_line_is_question:
            question = re.sub(r"^(Q|问|问题)[:：\s]*", "", lines[0], flags=re.I).strip()
            answer_lines = [re.sub(r"^(A|答|答案)[:：\s]*", "", line, flags=re.I).strip() for line in lines[1:]]
            answer = "\n".join(answer_lines).strip()
            if question and answer:
                add_entry(category_for_text(question + answer, block.source), question, answer, block.source)
                continue

        if "？" in text or "?" in text:
            parts = re.split(r"(?<=[？?])", text, maxsplit=1)
            if len(parts) == 2 and normalize_text(parts[0]) and normalize_text(parts[1]):
                question = normalize_text(parts[0])
                answer = normalize_text(parts[1])
                if len(question) <= 120 and "。" not in question:
                    add_entry(category_for_text(text, block.source), question, answer, block.source)

        match = re.match(r"^([^：:]{2,28})[:：]\s*(.+)$", text, flags=re.S)
        if match:
            label = normalize_text(match.group(1))
            answer = normalize_text(match.group(2))
            if is_qa_like_label(label) and len(answer) >= 8:
                add_entry(category_for_text(label + answer, block.source), make_question_from_label(label), answer, block.source, weight=15)
    return entries


def category_for_text(text: str, source: str) -> str:
    if any(word in text for word in ["邮箱", "邮件", "链接"]):
        return "邮件与链接"
    if any(word in text for word in ["退货", "退款", "换货", "售后"]):
        return "售后处理"
    if any(word in text for word in ["发货", "物流", "快递"]):
        return "发货物流"
    if any(word in text for word in ["价格", "优惠", "下单", "购买"]):
        return "售前转化"
    if any(word in source for word in ["极霸猫", "宅小兔"]):
        return "角色专属问答"
    return "通用咨询"


def is_qa_like_label(label: str) -> bool:
    if (
        len(label) > 18
        or "http" in label.lower()
        or any(char in label for char in ["，", "。", "（", "）", "(", ")", " "])
        or label.startswith("在推荐")
        or label.startswith("销售")
    ):
        return False
    return any(
        word in label
        for word in [
            "退款",
            "退货",
            "换货",
            "售后",
            "破损",
            "瑕疵",
            "发错",
            "丢件",
            "快递",
            "物流",
            "发货",
            "邮件",
            "邮箱",
            "图库",
            "定制",
            "补差价",
            "义乳",
            "义臀",
            "枕芯",
            "枕套",
            "抱枕",
        ]
    )


def make_question_from_label(label: str) -> str:
    if any(word in label for word in ["退款", "退货", "换货", "售后", "破损", "瑕疵", "发错", "丢件"]):
        return f"客户遇到{label}时怎么处理？"
    if any(word in label for word in ["邮件", "邮箱"]):
        return f"客户咨询{label}时怎么回复？"
    if any(word in label for word in ["图库", "定制", "补差价"]):
        return f"客户询问{label}时怎么说明？"
    return f"客户咨询{label}时怎么回复？"


def grouped_sources(blocks: list[Block]) -> list[str]:
    return sorted({block.source for block in blocks})


def add_raw_material_sections(doc: Document, blocks: list[Block], heading_level: int = 2):
    current_source = ""
    for block in blocks:
        if block.source != current_source:
            current_source = block.source
            doc.add_heading(current_source, level=heading_level)
        if "\n" in block.text and " | " in block.text:
            for line in block.text.splitlines():
                p = doc.add_paragraph()
                p.add_run(line)
        else:
            doc.add_paragraph(block.text)


def build_qa_doc(blocks: list[Block], out_path: Path):
    doc = Document()
    configure_doc(doc, "QA 问答知识库整理稿", "用于管理后台 QA 问答知识库录入：分类、问题、关键词、答案、权重、启用状态")
    add_source_note(doc, grouped_sources(blocks))
    add_toc(
        doc,
        [
            ("问答库名称", "建议按店铺/角色拆分，如「极霸猫售前售后 QA」「宅小兔售前售后 QA」「邮件与链接固定话术 QA」。"),
            ("分类", "按售前转化、售后处理、发货物流、邮件与链接、角色专属问答等场景归类。"),
            ("问题/关键词", "保留原文中的客户表达，并补充可用于关键词包含匹配的词。"),
            ("答案", "保留原文原意；只做换行、分段和场景归类。"),
        ],
    )
    entries = split_question_answer(blocks)
    doc.add_heading("可录入 QA 条目", level=1)
    if entries:
        grouped: dict[str, list[dict[str, str | list[str] | int | bool]]] = {}
        for entry in entries:
            grouped.setdefault(str(entry["category"]), []).append(entry)
        for category, category_entries in grouped.items():
            doc.add_heading(category, level=2)
            for index, entry in enumerate(category_entries, start=1):
                title = doc.add_paragraph()
                title.paragraph_format.space_before = Pt(4)
                title.paragraph_format.space_after = Pt(2)
                run = title.add_run(f"{index}. {entry['question']}")
                set_font(run, size=10.5, bold=True)
                table = doc.add_table(rows=0, cols=2)
                set_table_widths(table, [1320, 8040])
                for label, value in [
                    ("分类", entry["category"]),
                    ("问题", entry["question"]),
                    ("关键词", "、".join(entry["keywords"]) if entry["keywords"] else ""),
                    ("答案", entry["answer"]),
                    ("图片", entry["image_url"]),
                    ("权重", entry["weight"]),
                    ("启用", "是" if entry["enabled"] else "否"),
                    ("来源", entry["source"]),
                ]:
                    cells = table.add_row().cells
                    cells[0].text = str(label)
                    cells[1].text = str(value)
                format_table(table)
    else:
        doc.add_paragraph("未识别到标准问答格式，以下保留原文素材，建议人工按后台字段拆分。")

    doc.add_heading("待人工拆分的原文素材", level=1)
    add_raw_material_sections(doc, blocks)
    doc.save(out_path)
    return entries


def write_qa_csv(entries: list[dict[str, str | list[str] | int | bool]], out_path: Path):
    with out_path.open("w", encoding="utf-8-sig", newline="") as f:
        writer = csv.DictWriter(
            f,
            fieldnames=["category", "question", "keywords", "answer", "image_url", "weight", "enabled", "source"],
        )
        writer.writeheader()
        for entry in entries:
            writer.writerow(
                {
                    "category": entry["category"],
                    "question": entry["question"],
                    "keywords": "、".join(entry["keywords"]) if entry["keywords"] else "",
                    "answer": entry["answer"],
                    "image_url": entry["image_url"],
                    "weight": entry["weight"],
                    "enabled": "true" if entry["enabled"] else "false",
                    "source": entry["source"],
                }
            )


def write_usage_notes(summary: dict[str, object], out_path: Path):
    lines = [
        "# 测试知识库整理说明",
        "",
        "## 输出文件",
        "",
        "- `01-QA问答知识库整理稿.docx`：按后台 QA 字段整理的人工核对稿。",
        "- `01-QA问答知识库条目.csv`：可用于后续批量导入的 QA 候选条目，字段对应 `category/question/keywords/answer/image_url/weight/enabled`。",
        "- `02-产品知识库整理稿.docx`：按产品事实、售前策略、售后规则、邮件与链接规则等主题切片，可导入产品知识库。",
        "- `03-语气知识库整理稿.docx`：按人设、表达规范、示例语料整理，可用于语气知识库 persona 和语料参考。",
        "",
        "## 后台使用建议",
        "",
        "- QA 库：先导入 CSV 或按 DOCX 条目人工录入；命中后直接返回，不再检索产品库。",
        "- 产品库：优先导入客户可见的产品事实、售后规则和邮件流程；内部经营策略与明确“不能说”的内容建议单独建库或仅保留给人工客服参考。",
        "- 语气库：将整理稿中的“建议录入的虚拟人设”压缩为后台 `persona` 字段，再保留示例语料给 Prompt 或人工校准使用。",
        "",
        "## 覆盖情况",
        "",
        f"- 原始文档数：{len(summary['source_files'])}",
        f"- 抽取素材块：{summary['source_block_count']}",
        f"- QA 候选条目：{summary['qa_entry_count']}",
    ]
    out_path.write_text("\n".join(lines) + "\n", encoding="utf-8-sig")


def build_product_doc(blocks: list[Block], out_path: Path):
    doc = Document()
    configure_doc(doc, "产品知识库整理稿", "用于管理后台产品知识库导入与检索切片：产品/角色信息、售前售后规则、邮件链路约束")
    add_source_note(doc, grouped_sources(blocks))
    add_toc(
        doc,
        [
            ("产品知识库名称", "建议建立「角色售前售后产品知识库」和「销售策略与邮件规则知识库」。"),
            ("文档内容", "按可检索主题分段，每段只承载一个产品事实、规则或流程。"),
            ("检索使用", "普通咨询时由 AI 回复服务检索；QA 命中时不检索产品文档。"),
            ("禁说边界", "含「不能说」或内部经营策略的内容已单独归入内部策略段，不建议作为客户可见回复直接输出。"),
        ],
    )
    warning = doc.add_paragraph()
    warning.paragraph_format.space_before = Pt(4)
    warning.paragraph_format.space_after = Pt(10)
    r = warning.add_run("使用提醒：")
    set_font(r, bold=True)
    warning.add_run("产品库会被普通咨询检索调用，带有内部经营策略或明确“不能说”的素材应只作为内部理解材料，导入时建议单独建库并在机器人 Prompt 中限制不得直接对客户输出。")
    buckets = {
        "内部策略与禁说边界（不建议作为客户可见检索答案）": [],
        "角色与商品定位": [],
        "售前销售策略": [],
        "售后与履约规则": [],
        "邮件与链接规则": [],
        "其他可检索素材": [],
    }
    for block in blocks:
        text = block.text
        if "不能说" in block.source or any(word in text for word in ["不能出现在客户", "不能说", "客户语客户", "产品盈利模式"]):
            buckets["内部策略与禁说边界（不建议作为客户可见检索答案）"].append(block)
        elif any(word in text for word in ["极霸猫", "宅小兔", "角色", "性格"]):
            buckets["角色与商品定位"].append(block)
        elif any(word in text for word in ["销售", "策略", "优惠", "下单", "购买", "价格"]):
            buckets["售前销售策略"].append(block)
        elif any(word in text for word in ["售后", "退款", "退货", "换货", "物流", "发货"]):
            buckets["售后与履约规则"].append(block)
        elif any(word in text for word in ["邮件", "邮箱", "链接"]):
            buckets["邮件与链接规则"].append(block)
        else:
            buckets["其他可检索素材"].append(block)

    for heading, grouped in buckets.items():
        doc.add_heading(heading, level=1)
        if not grouped:
            doc.add_paragraph("暂无原文素材。")
            continue
        add_raw_material_sections(doc, grouped, heading_level=2)
    doc.save(out_path)


def build_tone_doc(blocks: list[Block], out_path: Path):
    doc = Document()
    configure_doc(doc, "语气知识库整理稿", "用于管理后台语气知识库录入：知识库名称、虚拟人设、表达边界、示例语料")
    add_source_note(doc, grouped_sources(blocks))
    add_toc(
        doc,
        [
            ("知识库名称", "建议按角色/风格拆分，如「二次元亲切语气库」「极霸猫角色语气库」「宅小兔角色语气库」。"),
            ("虚拟人设", "把性格特征、称呼、自称、回答长度和禁忌表达合并成一段可直接粘贴的 persona。"),
            ("示例语料", "保留原始语录和话术示例，供语气迁移和回复生成参考。"),
        ],
    )

    persona_blocks = [b for b in blocks if any(word in b.text for word in ["性格", "语气", "人设", "规范", "二次元", "称呼", "自称"])]
    corpus_blocks = [b for b in blocks if b not in persona_blocks]

    doc.add_heading("建议录入的虚拟人设", level=1)
    if persona_blocks:
        doc.add_paragraph("以下内容由原文中的性格特征、语言规范、二次元表达和角色话术整理而来，原意不变，可作为语气知识库 persona 初稿。")
        for block in persona_blocks:
            p = doc.add_paragraph()
            p.add_run(f"[{block.source}] ").bold = True
            p.add_run(block.text)
    else:
        doc.add_paragraph("未识别到明确人设描述。")

    doc.add_heading("示例语料与表达参考", level=1)
    add_raw_material_sections(doc, corpus_blocks or blocks)
    doc.save(out_path)


def write_extracts(blocks_by_file: dict[str, list[Block]]):
    EXTRACT_DIR.mkdir(parents=True, exist_ok=True)
    manifest = []
    for name, blocks in blocks_by_file.items():
        out = EXTRACT_DIR / (Path(name).stem + ".txt")
        text = "\n\n".join(f"[{block.kind}]\n{block.text}" for block in blocks)
        out.write_text(text, encoding="utf-8-sig")
        manifest.append({"source": name, "blocks": len(blocks), "chars": len(text), "extract": str(out)})
    (EXTRACT_DIR / "manifest.json").write_text(json.dumps(manifest, ensure_ascii=False, indent=2), encoding="utf-8-sig")


def main():
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)
    paths = sorted(SOURCE_DIR.glob("*.docx"))
    blocks_by_file = {path.name: list(iter_doc_blocks(path)) for path in paths}
    write_extracts(blocks_by_file)

    all_blocks = [block for blocks in blocks_by_file.values() for block in blocks]
    qa_blocks = [b for b in all_blocks if any(word in b.text or word in b.source for word in ["问", "答", "售前", "售后", "邮件", "邮箱", "链接", "发货", "退货", "退款", "优惠"])]
    product_blocks = [b for b in all_blocks if any(word in b.text or word in b.source for word in ["售前", "售后", "销售", "策略", "邮件", "极霸猫", "宅小兔", "产品", "物流", "发货", "退货", "退款"])]
    tone_blocks = [b for b in all_blocks if any(word in b.text or word in b.source for word in ["语气", "语料", "性格", "规范", "二次元", "极霸猫", "宅小兔"])]

    qa_entries = build_qa_doc(qa_blocks or all_blocks, OUTPUT_DIR / "01-QA问答知识库整理稿.docx")
    write_qa_csv(qa_entries, OUTPUT_DIR / "01-QA问答知识库条目.csv")
    build_product_doc(product_blocks or all_blocks, OUTPUT_DIR / "02-产品知识库整理稿.docx")
    build_tone_doc(tone_blocks or all_blocks, OUTPUT_DIR / "03-语气知识库整理稿.docx")

    summary = {
        "source_files": [path.name for path in paths],
        "source_block_count": len(all_blocks),
        "qa_block_count": len(qa_blocks),
        "product_block_count": len(product_blocks),
        "tone_block_count": len(tone_blocks),
        "qa_entry_count": len(qa_entries),
        "outputs": [
            str(OUTPUT_DIR / "01-QA问答知识库整理稿.docx"),
            str(OUTPUT_DIR / "01-QA问答知识库条目.csv"),
            str(OUTPUT_DIR / "02-产品知识库整理稿.docx"),
            str(OUTPUT_DIR / "03-语气知识库整理稿.docx"),
            str(OUTPUT_DIR / "整理说明.md"),
        ],
    }
    write_usage_notes(summary, OUTPUT_DIR / "整理说明.md")
    (OUTPUT_DIR / "整理说明.json").write_text(json.dumps(summary, ensure_ascii=False, indent=2), encoding="utf-8-sig")
    print(json.dumps(summary, ensure_ascii=False, indent=2))


if __name__ == "__main__":
    main()
